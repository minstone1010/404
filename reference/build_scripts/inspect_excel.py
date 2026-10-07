from pathlib import Path
import openpyxl
for p in Path('data').glob('*.xlsx'):
 w=openpyxl.load_workbook(p,read_only=True,data_only=True); print(p.name)
 for sh in w:
  print(sh.title,sh.max_row,sh.max_column,list(next(sh.values))[:40])
  if sh.title in ['channel_master','routing_rules','department_master']: print(list(sh.values)[1:4])
