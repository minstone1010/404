"""Portable local desktop and XLSX CLI. No server, telemetry or remote inference."""
import argparse,json,os,threading
from pathlib import Path
os.environ['HF_HUB_OFFLINE']='1';os.environ['TRANSFORMERS_OFFLINE']='1';os.environ['HF_HUB_DISABLE_TELEMETRY']='1'
from variants import Original,Office,Performance,result
ROOT=Path(__file__).resolve().parent

def load_engine():
    settings=json.loads((ROOT/'profile.json').read_text(encoding='utf-8'))
    model=ROOT/settings['model']
    if settings['profile']=='office':return Office(model,threads=settings.get('threads',4),model_min_score=settings.get('model_min_score',0.))
    if settings['profile']=='original':return Original(str(model),'cpu')
    import torch
    if not torch.cuda.is_available():raise RuntimeError('성능 중심 버전은 CUDA GPU가 필요합니다. 일반 CPU PC에서는 사무용 버전을 실행하세요.')
    return Performance(model,ROOT/settings['llm'],device='cuda',llm_batch=settings.get('llm_batch',8),support_min=settings['llm_support_min'],model_min_score=settings.get('model_min_score',0.),mask_types=settings.get('mask_types'),ner_dtype=settings.get('ner_dtype','float32'))

def read_rows(path,sheet=None,column='full_text'):
    import openpyxl
    w=openpyxl.load_workbook(path,read_only=True,data_only=True)
    if sheet=='원본(DB형식)' and sheet not in w.sheetnames and '시스템 양식' in w.sheetnames:
        sheet='시스템 양식'
    selected=w[sheet] if sheet else w[w.sheetnames[0]]
    it=selected.iter_rows(values_only=True);headers=list(next(it))
    if column=='full_text' and column not in headers and all(c in headers for c in ['접수번호','제목','질문','답변','처리자']):
        # Complaint model and schema-aware answer processing have separate scopes.
        rows=[]
        for values in it:
            ident=str(values[headers.index('접수번호')])
            fields={field:str(values[headers.index(field)] or '') for field in ['제목','질문','답변','처리자']}
            text=fields['제목']+'\n'+fields['질문']
            rows.append({'id':ident,'text':text,'fields':fields})
        w.close();return rows
    if column not in headers:raise ValueError(f'입력 열 {column!r}이 없습니다. --column 또는 화면에서 실제 열 이름을 지정하세요.')
    index=headers.index(column);id_index=headers.index('id') if 'id' in headers else None
    rows=[]
    for i,values in enumerate(it,1):rows.append({'id':str(values[id_index]) if id_index is not None else str(i),'text':'' if values[index] is None else str(values[index])})
    w.close();return rows

def export_rows(rows,path,include_original=False):
    import subprocess,tempfile,shutil
    node=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe'
    builder=next((parent/'masking_v3/report_tools/build_three_report.mjs' for parent in [ROOT,*ROOT.parents] if (parent/'masking_v3/report_tools/build_three_report.mjs').exists()),None)
    if node.exists() and builder:
        payload={'mode':'runtime','include_original':include_original,'rows':[{k:v for k,v in row.items() if include_original or k!='original'} for row in rows]}
        with tempfile.TemporaryDirectory(dir=Path(path).resolve().parent) as temp:
            source=Path(temp)/'payload.json';source.write_text(json.dumps(payload,ensure_ascii=False),encoding='utf-8')
            subprocess.run([str(node),str(builder),'--runtime',str(source),str(Path(path).resolve())],check=True)
        return
    # Portable fallback: the proprietary Artifact Tool is absent on ordinary offline office PCs.
    # Authored comparison deliverables are generated with Artifact Tool, not this fallback.
    import openpyxl
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.worksheet.table import Table,TableStyleInfo
    w=openpyxl.Workbook();s=w.active;s.title='마스킹 결과'
    headers=['문서 ID','마스킹 결과','탐지 수','탐지 유형']
    if include_original:headers.append('원문 (검토용)')
    s.append(headers)
    for row in rows:
        values=[row['id'],row['masked_text'],len(row['spans']),', '.join(sorted({x['type'] for x in row['spans']}))]
        if include_original:values.append(row['original'])
        for col,value in enumerate(values,1):
            cell=s.cell(s.max_row+1 if col==1 else s.max_row,col,value)
            if isinstance(value,str):cell.data_type='s'
    for cell in s[1]:cell.font=Font(name='맑은 고딕',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='233B63')
    import math,unicodedata
    for row in s.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
        widths={1:18,2:85,3:12,4:35,5:85}
        lines=1
        for cell in row:
            text='' if cell.value is None else str(cell.value)
            needed=sum(max(1,math.ceil(sum(2 if unicodedata.east_asian_width(ch) in {'W','F'} else 1 for ch in line)/(widths[cell.column]-2))) for line in text.split('\n'))
            lines=max(lines,needed)
        s.row_dimensions[row[0].row].height=min(409,14*lines+4)
    s.column_dimensions['A'].width=18;s.column_dimensions['B'].width=85;s.column_dimensions['C'].width=12;s.column_dimensions['D'].width=35
    if include_original:s.column_dimensions['E'].width=85
    s.freeze_panes='B2';s.sheet_view.showGridLines=False
    if rows:
        from openpyxl.utils import get_column_letter
        table=Table(displayName='MaskingResults',ref=f'A1:{get_column_letter(len(headers))}{len(rows)+1}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);s.add_table(table)
    w.save(path)

def process_file(engine,input_path,output_path,sheet,column,include_original=False,progress=None):
    if Path(input_path).resolve()==Path(output_path).resolve():raise ValueError('출력 경로는 원본 파일과 달라야 합니다.')
    rows=read_rows(input_path,sheet,column)
    if hasattr(engine,'prepare'):engine.prepare([r['text'] for r in rows],progress)
    results=[]
    for i,row in enumerate(rows,1):
        answer=result(engine,row['text'])
        original=row['text']
        if 'fields' in row:
            from system_form import combined_answer
            answer=combined_answer(engine,row['fields'],answer)
            original+='\n\n[답변]\n'+row['fields']['답변']+'\n\n[처리자]\n'+row['fields']['처리자']
        answer.update(id=row['id'],original=original);results.append(answer)
        if progress and i%20==0:progress(i,len(rows))
    export_rows(results,output_path,include_original)
    if hasattr(engine,'llm_cache'):engine.llm_cache.clear()
    return len(rows)

def main():
    p=argparse.ArgumentParser();p.add_argument('--input');p.add_argument('--output');p.add_argument('--sheet');p.add_argument('--column',default='full_text');p.add_argument('--include-original',action='store_true');a=p.parse_args()
    if a.input:
        if not a.output:p.error('--output required with --input')
        count=process_file(load_engine(),a.input,a.output,a.sheet,a.column,a.include_original)
        print(f'{count}건 마스킹 결과 저장 완료: {a.output}');return
    import tkinter as tk
    from tkinter import filedialog,messagebox
    config=json.loads((ROOT/'profile.json').read_text(encoding='utf-8'))
    root=tk.Tk();root.title(config['name']+' — 로컬 개인정보 마스킹');root.geometry('950x820')
    tk.Label(root,text=config['name']).pack()
    tk.Label(root,text='마스킹 대상은 이름·주소·이메일·휴대폰번호·승차권번호 5종입니다.').pack()
    source=tk.Text(root,height=10);source.pack(fill='both',expand=True)
    tk.Label(root,text='마스킹 결과').pack();target=tk.Text(root,height=10);target.pack(fill='both',expand=True)
    status=tk.StringVar(value='모델은 첫 실행 때 로컬 파일에서 불러옵니다.');tk.Label(root,textvariable=status).pack()
    engine=[None];lock=threading.Lock()
    def task(fn):
        if not lock.acquire(blocking=False):return
        status.set('처리 중…')
        def worker():
            try:
                if engine[0] is None:engine[0]=load_engine()
                message=fn(engine[0]);root.after(0,lambda:status.set(message))
            except Exception as e:
                msg=str(e);root.after(0,lambda:messagebox.showerror('처리 실패',msg));root.after(0,lambda:status.set('처리 실패'))
            finally:lock.release()
        threading.Thread(target=worker,daemon=True).start()
    def text_run():
        text=source.get('1.0','end-1c')
        def work(e):
            answer=result(e,text)
            def show():target.delete('1.0','end');target.insert('1.0',answer['masked_text'])
            root.after(0,show)
            if hasattr(e,'llm_cache'):e.llm_cache.clear()
            return '마스킹 완료'
        task(work)
    tk.Button(root,text='입력 문장 마스킹',command=text_run).pack()
    panel=tk.Frame(root);panel.pack(fill='x')
    tk.Label(panel,text='엑셀 시트 이름 (비우면 첫 시트)').pack(side='left');sheet_box=tk.Entry(panel,width=23);sheet_box.insert(0,'원본(DB형식)');sheet_box.pack(side='left')
    tk.Label(panel,text='민원 텍스트 열').pack(side='left');column_box=tk.Entry(panel,width=20);column_box.insert(0,'full_text');column_box.pack(side='left')
    include=tk.BooleanVar(value=False);tk.Checkbutton(root,text='원문도 함께 저장 (검토용)',variable=include).pack()
    def file_run():
        src=filedialog.askopenfilename(filetypes=[('Excel workbook','*.xlsx')])
        if not src:return
        dest=filedialog.asksaveasfilename(defaultextension='.xlsx',initialfile='마스킹결과.xlsx',filetypes=[('Excel workbook','*.xlsx')])
        if not dest:return
        sheet=sheet_box.get().strip() or None;column=column_box.get().strip();keep=include.get()
        def work(e):
            n=process_file(e,src,dest,sheet,column,keep,lambda done,total:root.after(0,lambda:status.set(f'처리 중 {done}/{total}')))
            return f'{n}건 엑셀 저장 완료'
        task(work)
    tk.Button(root,text='엑셀 파일 선택 → 마스킹 결과 엑셀 저장',command=file_run).pack()
    def clear():
        if lock.locked():return
        source.delete('1.0','end');target.delete('1.0','end')
        if engine[0] is not None and hasattr(engine[0],'llm_cache'):engine[0].llm_cache.clear()
    tk.Button(root,text='입력·결과 지우기',command=clear).pack();root.mainloop()
if __name__=='__main__':main()
